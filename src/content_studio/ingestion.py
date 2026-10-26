"""Input acquisition: local files and URLs (PRD Phase 1).

External material is treated strictly as data. Nothing here executes scripts,
follows prompt instructions found in documents, or writes outside the project.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from bs4.element import NavigableString
from docx import Document as DocxDocument
from pypdf import PdfReader

from content_studio.errors import IngestionError
from content_studio.models import (
    AccessStatus,
    ProjectConfig,
    SourceKind,
    SourceMap,
    SourceRecord,
    SourceRegistry,
)
from content_studio.paths import require_inside
from content_studio.settings import Settings
from content_studio.source_map import build_document_map, load_source_map, write_source_map
from content_studio.utils import (
    read_text_utf8,
    relative_posix,
    sha256_bytes,
    slugify,
    utc_now,
    write_json,
    write_text,
)

_SUFFIX_KIND: dict[str, SourceKind] = {
    ".pdf": "pdf",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".text": "text",
    ".html": "html",
    ".htm": "html",
    ".docx": "docx",
}

_BLANK_LINES = re.compile(r"\n{3,}")


@dataclass
class LoadedDocument:
    kind: SourceKind
    title: str
    text: str
    origin: str
    raw: bytes = b""
    suffix: str = ".txt"
    access_status: AccessStatus = "fetched"
    notes: str = ""
    publisher: str | None = None
    published_at: str | None = None
    # ``(page_number, normalized_page_text)`` for PDFs; empty pages keep their
    # number so page locators stay truthful.
    pages: list[tuple[int, str]] | None = None
    extra_fields: dict[str, str] = field(default_factory=dict)


def _normalize(text: str) -> str:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = "\n".join(line.rstrip() for line in cleaned.split("\n"))
    return _BLANK_LINES.sub("\n\n", cleaned).strip()


def _title_from_text(text: str, fallback: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            candidate = stripped.lstrip("#").strip()
            if candidate:
                return candidate[:200]
        return stripped[:200]
    return fallback


def _extract_pdf(data: bytes) -> tuple[str, AccessStatus, str, list[tuple[int, str]]]:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises several unrelated exception types
        return "", "unavailable", f"PDF could not be opened: {exc.__class__.__name__}", []
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            return "", "unavailable", "PDF is encrypted", []
    pages: list[tuple[int, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:
            raw = ""
        pages.append((index, _normalize(raw)))
    # Normalise per page so page offsets survive; the "\n\n" join is what
    # ``_page_spans`` reconstructs when it builds the source map.
    text = "\n\n".join(page_text for _, page_text in pages if page_text)
    if not text:
        return (
            "",
            "partial",
            "PDF has no extractable text; OCR is not implemented (Phase 3)",
            pages,
        )
    return text, "fetched", "", pages


def _extract_docx(data: bytes) -> tuple[str, AccessStatus, str]:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        return "", "unavailable", f"DOCX could not be opened: {exc.__class__.__name__}"
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = _normalize("\n".join(parts))
    if not text:
        return "", "partial", "DOCX contains no extractable paragraph text"
    return text, "fetched", ""


def _html_parts(data: bytes) -> tuple[str, str]:
    soup = BeautifulSoup(data, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    _collapse_math(soup)
    text = _normalize(soup.get_text("\n", strip=True))
    return title, text


def _collapse_math(soup: BeautifulSoup) -> None:
    """Replace each ``<math>`` subtree with one readable inline token.

    LaTeXML (which arXiv HTML uses) splits a formula into ``<mi>``/``<mn>``
    nodes *and* repeats it as an ``<annotation encoding="application/x-tex">``.
    Calling ``get_text`` on that tree emits one character per line, which is
    how evidence excerpts ended up as ``(\\nx\\n1\\n,\\n…``.

    The TeX annotation is the clean, single-line form, so it is preferred;
    without one the visible tokens are joined with single spaces. Replacing
    the whole subtree also removes the duplicate rendering.
    """

    for math in soup.find_all("math"):
        annotation = math.find("annotation")
        if annotation is not None:
            rendered = annotation.get_text(" ", strip=True)
        else:
            rendered = math.get_text(" ", strip=True)
        if not rendered:
            math.decompose()
            continue
        math.replace_with(NavigableString(f" {rendered} "))


def load_local(path: Path, settings: Settings) -> LoadedDocument:
    if not path.is_file():
        raise IngestionError(f"input file not found: {path}")
    size = path.stat().st_size
    if size > settings.max_source_bytes:
        raise IngestionError(f"input file exceeds {settings.max_source_bytes} bytes: {path}")
    data = path.read_bytes()
    suffix = path.suffix.lower()
    kind = _SUFFIX_KIND.get(suffix)
    if kind is None:
        raise IngestionError(f"unsupported input type '{suffix or path.name}': {path}")

    notes = ""
    access: AccessStatus = "fetched"
    pages: list[tuple[int, str]] | None = None
    if kind == "pdf":
        text, access, notes, pages = _extract_pdf(data)
        title = path.stem
    elif kind == "docx":
        text, access, notes = _extract_docx(data)
        title = path.stem
    elif kind == "html":
        html_title, text = _html_parts(data)
        title = html_title or path.stem
    else:
        text = _normalize(read_text_utf8(path))
        title = _title_from_text(text, path.stem)

    if not text and access == "fetched":
        access = "partial"
        notes = notes or "no text could be extracted"

    return LoadedDocument(
        kind=kind,
        title=title or path.stem,
        text=text,
        origin=str(path),
        raw=data,
        suffix=suffix,
        access_status=access,
        notes=notes,
        pages=pages,
    )


def _kind_from_content_type(content_type: str, url: str) -> SourceKind:
    lowered = content_type.lower()
    if "pdf" in lowered or urlparse(url).path.lower().endswith(".pdf"):
        return "pdf"
    if "html" in lowered or "xml" in lowered:
        return "html"
    if "markdown" in lowered:
        return "markdown"
    return "text"


def fetch_url(url: str, settings: Settings) -> LoadedDocument:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise IngestionError(f"only http(s) URLs are supported: {url}")
    try:
        with httpx.Client(
            follow_redirects=True,
            timeout=settings.request_timeout_seconds,
            headers={"User-Agent": "LCES/0.2 (+local content research)"},
        ) as client:
            response = client.get(url)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        return LoadedDocument(
            kind="web",
            title=parsed.netloc or url,
            text="",
            origin=url,
            access_status="unavailable",
            notes=f"fetch failed: {exc.__class__.__name__}",
        )

    data = response.content
    if len(data) > settings.max_source_bytes:
        raise IngestionError(f"remote source exceeds {settings.max_source_bytes} bytes: {url}")

    content_type = response.headers.get("content-type", "")
    kind = _kind_from_content_type(content_type, url)
    publisher = parsed.netloc
    notes = ""

    pages: list[tuple[int, str]] | None = None
    if kind == "pdf":
        text, access, notes, pages = _extract_pdf(data)
        title = Path(parsed.path).stem or publisher
        suffix = ".pdf"
    elif kind == "html":
        html_title, text = _html_parts(data)
        title = html_title or publisher
        access = "fetched" if text else "partial"
        notes = notes or ("no text could be extracted from HTML" if not text else "")
        suffix = ".html"
    else:
        text = _normalize(response.text)
        title = _title_from_text(text, publisher)
        access = "fetched" if text else "partial"
        suffix = ".txt"

    return LoadedDocument(
        kind=kind,
        title=title,
        text=text,
        origin=url,
        raw=data,
        suffix=suffix,
        access_status=access,
        notes=notes,
        publisher=publisher,
        pages=pages,
    )


def _next_source_id(existing: list[SourceRecord]) -> str:
    highest = 0
    for record in existing:
        try:
            highest = max(highest, int(record.source_id.split("-")[1]))
        except (IndexError, ValueError):
            continue
    return f"source-{highest + 1:03d}"


def load_registry(project_dir: Path) -> SourceRegistry | None:
    path = project_dir / "internal" / "source_registry.json"
    if not path.exists():
        return None
    try:
        return SourceRegistry.model_validate_json(path.read_text(encoding="utf-8"))
    except ValueError as exc:
        raise IngestionError(f"invalid source registry: {path}: {exc}") from exc


def save_registry(project_dir: Path, registry: SourceRegistry) -> Path:
    path = project_dir / "internal" / "source_registry.json"
    write_json(path, registry)
    return path


def _store_document(
    project_dir: Path,
    config: ProjectConfig,
    document: LoadedDocument,
    source_id: str,
) -> SourceRecord:
    raw_name = f"{source_id}-{slugify(document.title)}{document.suffix}"
    local_path = project_dir / "input" / raw_name
    if document.raw:
        local_path.parent.mkdir(parents=True, exist_ok=True)
        local_path.write_bytes(document.raw)

    text_path = project_dir / "internal" / "sources" / f"{source_id}.txt"
    stored_text = document.text or "[NO EXTRACTED TEXT]"
    write_text(text_path, stored_text)

    # The map is written next to the source text at ingest time: PDF page
    # numbers are only available here, and later stages must not guess them.
    document_map = build_document_map(
        source_id,
        stored_text,
        pages=document.pages,
        sha256=sha256_bytes(document.raw) if document.raw else "",
        parser=f"{document.kind}/v1",
    )
    existing = load_source_map(project_dir)
    documents = (
        [item for item in existing.documents if item.source_id != source_id]
        if existing is not None
        else []
    )
    documents.append(document_map)
    write_source_map(
        project_dir,
        SourceMap(project_id=config.project_id, documents=documents),
    )

    return SourceRecord(
        source_id=source_id,
        title=document.title,
        kind=document.kind,
        origin=document.origin,
        local_path=local_path.name if document.raw else None,
        retrieved_at=utc_now(),
        sha256=sha256_bytes(document.raw) if document.raw else None,
        size_bytes=len(document.raw) if document.raw else None,
        publisher=document.publisher,
        published_at=document.published_at,
        access_status=document.access_status,
        text_chars=len(document.text),
        notes=document.notes or None,
    )


def ingest_project(
    project_dir: Path,
    config: ProjectConfig,
    settings: Settings,
    *,
    extra_paths: list[str] | None = None,
    extra_urls: list[str] | None = None,
    allow_external: bool = False,
) -> SourceRegistry:
    root = project_dir.resolve()
    registry = load_registry(root) or SourceRegistry(project_id=config.project_id)

    paths = [*config.input_paths, *(extra_paths or [])]
    urls = [*config.source_urls, *(extra_urls or [])]

    if not paths and not urls and not registry.sources:
        raise IngestionError(
            "no inputs configured; add input_paths/source_urls to project.yaml or pass --path/--url"
        )

    for raw_path in paths:
        candidate = Path(raw_path)
        if not candidate.is_absolute():
            candidate = root / candidate
        if not allow_external:
            require_inside(candidate, root, what="input path")
        elif not candidate.is_file():
            raise IngestionError(f"input file not found: {candidate}")
        document = load_local(candidate, settings)
        document.origin = (
            relative_posix(candidate, root) if _is_relative(candidate, root) else str(candidate)
        )
        source_id = _next_source_id(registry.sources)
        registry.sources.append(_store_document(root, config, document, source_id))

    for url in urls:
        document = fetch_url(url, settings)
        source_id = _next_source_id(registry.sources)
        registry.sources.append(_store_document(root, config, document, source_id))

    save_registry(root, registry)
    return registry


def _is_relative(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return True


def load_source_text(project_dir: Path, source_id: str) -> str:
    path = project_dir / "internal" / "sources" / f"{source_id}.txt"
    if not path.exists():
        return ""
    return read_text_utf8(path)
